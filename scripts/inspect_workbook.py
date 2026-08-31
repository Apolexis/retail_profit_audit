"""Read-only inventory of the uploaded retail workbook.

This script does not edit the workbook. It produces a JSON model map used by
the store-level extraction and formula-audit passes.
"""

import json
import os
import re
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path("/home/ubuntu/upload/учет2026.xlsm")
OUT_DIR = Path("/home/ubuntu/retail_profit_audit_2026/audit_work")
OUT_DIR.mkdir(parents=True, exist_ok=True)


def safe_value(value):
    if isinstance(value, datetime):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def cell_preview(ws, max_rows=80, max_cols=25):
    preview = []
    for row in ws.iter_rows(
        min_row=1,
        max_row=min(ws.max_row, max_rows),
        min_col=1,
        max_col=min(ws.max_column, max_cols),
    ):
        values = []
        for cell in row:
            value = cell.value
            if value is not None:
                values.append({"cell": cell.coordinate, "value": safe_value(value)})
        if values:
            preview.append(values)
    return preview


def find_formula_errors(ws):
    error_cells = []
    for row in ws.iter_rows():
        for cell in row:
            value = cell.value
            if isinstance(value, str) and value.startswith("#"):
                error_cells.append({"cell": cell.coordinate, "value": value})
    return error_cells


def main():
    wb = load_workbook(SOURCE, read_only=False, data_only=False, keep_vba=True)
    model = {
        "source": str(SOURCE),
        "generated_at": datetime.now().isoformat(),
        "workbook": {
            "sheet_count": len(wb.worksheets),
            "sheet_names": wb.sheetnames,
            "defined_names": [
                {
                    "name": name,
                    "value": str(defined_name.attr_text),
                }
                for name, defined_name in wb.defined_names.items()
            ],
            "calculation": {
                "fullCalcOnLoad": getattr(wb.calculation, "fullCalcOnLoad", None),
                "forceFullCalc": getattr(wb.calculation, "forceFullCalc", None),
                "iterate": getattr(wb.calculation, "iterate", None),
            },
        },
        "sheets": [],
    }

    for position, ws in enumerate(wb.worksheets, start=1):
        formulas = []
        nonempty_count = 0
        hidden_rows = []
        hidden_columns = []
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    nonempty_count += 1
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas.append({"cell": cell.coordinate, "formula": cell.value})
        for index, dimension in ws.row_dimensions.items():
            if dimension.hidden:
                hidden_rows.append(index)
        for letter, dimension in ws.column_dimensions.items():
            if dimension.hidden:
                hidden_columns.append(letter)
        model["sheets"].append(
            {
                "position": position,
                "name": ws.title,
                "state": ws.sheet_state,
                "dimensions": {"max_row": ws.max_row, "max_column": ws.max_column},
                "nonempty_cells": nonempty_count,
                "merged_ranges": [str(merged) for merged in ws.merged_cells.ranges],
                "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
                "hidden_rows": hidden_rows,
                "hidden_columns": hidden_columns,
                "formula_count": len(formulas),
                "formula_samples": formulas[:50],
                "formula_errors": find_formula_errors(ws),
                "print_area": str(ws.print_area) if ws.print_area else None,
                "preview": cell_preview(ws),
            }
        )

    output = OUT_DIR / "inventory.json"
    output.write_text(json.dumps(model, ensure_ascii=False, indent=2), encoding="utf-8")
    print(output)
    print(f"Sheets: {len(wb.worksheets)}")
    for sheet in model["sheets"]:
        print(
            f"{sheet['position']:>2}. {sheet['name']} | {sheet['state']} | "
            f"{sheet['dimensions']['max_row']}x{sheet['dimensions']['max_column']} | "
            f"formulas: {sheet['formula_count']}"
        )


if __name__ == "__main__":
    main()
