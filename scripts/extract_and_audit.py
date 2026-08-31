"""Extract store-level P&L data and audit formulas from the uploaded workbook.

The source workbook is opened read-only in practice: it is never saved or edited.
All checks are performed against cached values and formula strings, then saved to
CSV/JSON for downstream reporting.
"""

import csv
import json
import math
import re
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


SOURCE = Path("/home/ubuntu/upload/учет2026.xlsm")
OUT_DIR = Path("/home/ubuntu/retail_profit_audit_2026/audit_work")
OUT_DIR.mkdir(parents=True, exist_ok=True)

MONTHS = [
    "Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
    "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь",
]
METRIC_COLUMNS = {
    "purchases": "G",
    "revenue": "H",
    "gross_profit": "L",
    "cash_revenue": "M",
    "cashless_revenue": "N",
    "receipts_total": "O",
    "writeoffs_copy": "P",
    "writeoffs_frozen": "Q",
    "movement": "R",
    "discount": "S",
    "revaluation": "T",
    "cash_operating_costs": "AD",
    "driver_cash": "AE",
    "utilities_cash": "AF",
    "rent": "AG",
    "bank_fee": "AH",
    "gross_profit_tax": "AI",
    "salary_cashless": "AJ",
    "payroll_tax": "AK",
    "vacation_cashless": "AL",
    "vacation_tax": "AM",
    "salary_cash": "AN",
    "vacation_cash": "AO",
    "personal_income_tax": "AP",
}


def value_or_zero(value):
    if isinstance(value, bool) or value is None or isinstance(value, str):
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def is_formula(value):
    return isinstance(value, str) and value.startswith("=")


def normalized(formula):
    if not isinstance(formula, str):
        return ""
    return re.sub(r"\s+", "", formula).upper()


def close_enough(actual, expected, tolerance=0.15):
    return abs(value_or_zero(actual) - float(expected)) <= tolerance


def write_csv(path, rows, fieldnames):
    with path.open("w", newline="", encoding="utf-8-sig") as output:
        writer = csv.DictWriter(output, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def month_blocks(ws):
    headers = []
    for row in range(1, ws.max_row + 1):
        if ws.cell(row, 1).value in MONTHS:
            headers.append((row, ws.cell(row, 1).value))
    blocks = []
    for index, (start, month) in enumerate(headers):
        next_start = headers[index + 1][0] if index + 1 < len(headers) else ws.max_row + 1
        total_row = None
        for row in range(start + 1, next_start):
            if ws.cell(row, 1).value == "Итого":
                total_row = row
                break
        if total_row is None:
            continue
        data_rows = [
            row for row in range(start + 1, total_row)
            if isinstance(ws.cell(row, 1).value, (int, float))
        ]
        if data_rows:
            blocks.append(
                {
                    "month": month,
                    "start_row": start,
                    "total_row": total_row,
                    "data_rows": data_rows,
                }
            )
    return blocks


def issue(issues, sheet, cell, severity, category, description, impact=""):
    issues.append(
        {
            "sheet": sheet,
            "cell": cell,
            "severity": severity,
            "category": category,
            "description": description,
            "impact": impact,
        }
    )


def check_formula_presence(formulas_ws, sheet_name, coordinate, issues, expected_prefix=None):
    formula = formulas_ws[coordinate].value
    if not is_formula(formula):
        issue(
            issues,
            sheet_name,
            coordinate,
            "Warning",
            "Formula consistency",
            "Ожидалась расчетная формула, но в ячейке нет формулы.",
            "Требует проверки: возможна ручная замена расчета.",
        )
        return False
    if expected_prefix and not normalized(formula).startswith(normalized(expected_prefix)):
        issue(
            issues,
            sheet_name,
            coordinate,
            "Warning",
            "Formula consistency",
            f"Формула не соответствует ожидаемой логике {expected_prefix}.",
            f"Фактическая формула: {formula}",
        )
        return False
    return True


def audit_month(formulas_ws, values_ws, store, block, issues, prior_data_row):
    month = block["month"]
    total_row = block["total_row"]
    data_rows = block["data_rows"]
    first_row = data_rows[0]

    for row in data_rows:
        # External purchase/sale links should remain formulas when present.
        for col in ("C", "D", "E", "F"):
            formula = formulas_ws[f"{col}{row}"].value
            if not is_formula(formula):
                issue(
                    issues,
                    store,
                    f"{col}{row}",
                    "Info",
                    "Source link",
                    "Ячейка дневного движения не содержит формулу-ссылку на технический лист.",
                    "Возможна допустимая ручная корректировка; требуется подтвердить источник.",
                )
            elif "#REF!" in formula.upper():
                issue(
                    issues,
                    store,
                    f"{col}{row}",
                    "Critical",
                    "Broken reference",
                    "Формула содержит #REF!.",
                    "Исходный поток закупки/продажи разорван.",
                )

        check_formula_presence(formulas_ws, store, f"G{row}", issues, "=SUM(")
        check_formula_presence(formulas_ws, store, f"H{row}", issues, "=SUM(")
        check_formula_presence(formulas_ws, store, f"I{row}", issues, "=IFERROR(")
        check_formula_presence(formulas_ws, store, f"J{row}", issues, "=IFERROR(")
        check_formula_presence(formulas_ws, store, f"K{row}", issues, "=IFERROR(")
        check_formula_presence(formulas_ws, store, f"L{row}", issues, "=H")
        check_formula_presence(formulas_ws, store, f"O{row}", issues, "=M")
        check_formula_presence(formulas_ws, store, f"AD{row}", issues, "=SUM(")
        check_formula_presence(formulas_ws, store, f"AH{row}", issues, "=N")
        check_formula_presence(formulas_ws, store, f"AI{row}", issues, "=L")
        check_formula_presence(formulas_ws, store, f"AP{row}", issues, "=(AD")

        # Recalculate core daily relationships against cached values.
        daily_checks = {
            f"G{row}": value_or_zero(values_ws[f"C{row}"].value) + value_or_zero(values_ws[f"D{row}"].value),
            f"H{row}": value_or_zero(values_ws[f"E{row}"].value) + value_or_zero(values_ws[f"F{row}"].value),
            f"L{row}": value_or_zero(values_ws[f"H{row}"].value) - value_or_zero(values_ws[f"G{row}"].value),
            f"O{row}": value_or_zero(values_ws[f"M{row}"].value) + value_or_zero(values_ws[f"N{row}"].value),
            f"AD{row}": sum(value_or_zero(values_ws[f"{get_column_letter(col)}{row}"].value) for col in range(21, 30)),
            f"AH{row}": -value_or_zero(values_ws[f"N{row}"].value) * 0.02,
            f"AI{row}": -value_or_zero(values_ws[f"L{row}"].value) * 0.12,
            f"AP{row}": -(
                value_or_zero(values_ws[f"AD{row}"].value)
                + value_or_zero(values_ws[f"AN{row}"].value)
                + value_or_zero(values_ws[f"AO{row}"].value)
            ) * 0.22,
        }
        for coordinate, expected in daily_checks.items():
            actual = values_ws[coordinate].value
            if not close_enough(actual, expected):
                issue(
                    issues,
                    store,
                    coordinate,
                    "Critical",
                    "Formula value tie-out",
                    f"Кэшированное значение {value_or_zero(actual):,.2f} не сходится с независимым пересчетом {expected:,.2f}.",
                    f"Отклонение: {value_or_zero(actual) - expected:,.2f}.",
                )

        if prior_data_row is not None:
            formula = formulas_ws[f"B{row}"].value
            expected = (
                value_or_zero(values_ws[f"B{prior_data_row}"].value)
                + value_or_zero(values_ws[f"H{prior_data_row}"].value)
                - value_or_zero(values_ws[f"O{prior_data_row}"].value)
                - value_or_zero(values_ws[f"P{prior_data_row}"].value)
                - value_or_zero(values_ws[f"Q{prior_data_row}"].value)
                + value_or_zero(values_ws[f"R{prior_data_row}"].value)
                - value_or_zero(values_ws[f"S{prior_data_row}"].value)
                + value_or_zero(values_ws[f"T{prior_data_row}"].value)
            )
            if not is_formula(formula):
                issue(
                    issues,
                    store,
                    f"B{row}",
                    "Warning",
                    "Stock roll-forward",
                    "Ожидалась формула переходящего остатка, но обнаружено ручное значение.",
                    "Возможен разрыв цепочки остатков.",
                )
            if not close_enough(values_ws[f"B{row}"].value, expected):
                issue(
                    issues,
                    store,
                    f"B{row}",
                    "Critical",
                    "Stock roll-forward",
                    "Переходящий остаток не сходится с предыдущим днем.",
                    f"Отклонение: {value_or_zero(values_ws[f'B{row}'].value) - expected:,.2f}.",
                )
        prior_data_row = row

    # Monthly totals should tie to daily figures.
    for col in [
        "C", "D", "E", "F", "G", "H", "L", "M", "N", "O", "P", "Q", "R", "S", "T",
        "U", "V", "W", "X", "Y", "Z", "AA", "AB", "AC", "AD", "AE", "AF", "AG", "AH",
        "AI", "AJ", "AK", "AL", "AM", "AN", "AO", "AP",
    ]:
        coordinate = f"{col}{total_row}"
        expected = sum(value_or_zero(values_ws[f"{col}{row}"].value) for row in data_rows)
        if not close_enough(values_ws[coordinate].value, expected):
            issue(
                issues,
                store,
                coordinate,
                "Critical",
                "Monthly total",
                "Итог по месяцу не сходится с суммой дневных строк.",
                f"Отклонение: {value_or_zero(values_ws[coordinate].value) - expected:,.2f}.",
            )
    for col in ("I", "J", "K"):
        numeric_values = [
            values_ws[f"{col}{row}"].value for row in data_rows
            if isinstance(values_ws[f"{col}{row}"].value, (int, float))
        ]
        if numeric_values:
            expected = sum(numeric_values) / len(numeric_values)
            coordinate = f"{col}{total_row}"
            if not close_enough(values_ws[coordinate].value, expected):
                issue(
                    issues,
                    store,
                    coordinate,
                    "Critical",
                    "Monthly average",
                    "Средний процент месяца не сходится со средним дневных значений.",
                    f"Отклонение: {value_or_zero(values_ws[coordinate].value) - expected:,.2f} п.п.",
                )

    net_profit_coordinate = f"AQ{first_row}"
    net_expected = (
        value_or_zero(values_ws[f"L{total_row}"].value)
        - value_or_zero(values_ws[f"Q{total_row}"].value)
        - value_or_zero(values_ws[f"AD{total_row}"].value)
        - value_or_zero(values_ws[f"AE{total_row}"].value)
        - value_or_zero(values_ws[f"AF{total_row}"].value)
        - value_or_zero(values_ws[f"AG{total_row}"].value)
        + value_or_zero(values_ws[f"AH{total_row}"].value)
        + value_or_zero(values_ws[f"AI{total_row}"].value)
        - value_or_zero(values_ws[f"AJ{total_row}"].value)
        - value_or_zero(values_ws[f"AK{total_row}"].value)
        - value_or_zero(values_ws[f"AL{total_row}"].value)
        - value_or_zero(values_ws[f"AM{total_row}"].value)
        - value_or_zero(values_ws[f"AN{total_row}"].value)
        - value_or_zero(values_ws[f"AO{total_row}"].value)
        + value_or_zero(values_ws[f"AP{total_row}"].value)
    )
    if not close_enough(values_ws[net_profit_coordinate].value, net_expected):
        issue(
            issues,
            store,
            net_profit_coordinate,
            "Critical",
            "Net profit formula",
            "Месячная чистая прибыль не сходится с независимым пересчетом по используемым формулой статьям.",
            f"Отклонение: {value_or_zero(values_ws[net_profit_coordinate].value) - net_expected:,.2f}.",
        )
    formula = formulas_ws[net_profit_coordinate].value
    if not is_formula(formula):
        issue(
            issues,
            store,
            net_profit_coordinate,
            "Critical",
            "Net profit formula",
            "Месячная чистая прибыль не содержит расчетной формулы.",
            "Итоговый финансовый результат месяца не является аудируемым расчетом.",
        )

    metric_values = {
        metric: value_or_zero(values_ws[f"{col}{total_row}"].value)
        for metric, col in METRIC_COLUMNS.items()
    }
    metric_values.update(
        {
            "store": store,
            "month": month,
            "month_number": MONTHS.index(month) + 1,
            "total_row": total_row,
            "net_profit": value_or_zero(values_ws[net_profit_coordinate].value),
            "net_profit_cell": net_profit_coordinate,
        }
    )
    return metric_values, prior_data_row


def formulas_error_scan(formulas_ws, values_ws, sheet_name, issues):
    for row in formulas_ws.iter_rows():
        for formula_cell in row:
            formula = formula_cell.value
            cached = values_ws[formula_cell.coordinate].value
            if is_formula(formula) and "#REF!" in formula.upper():
                issue(
                    issues,
                    sheet_name,
                    formula_cell.coordinate,
                    "Critical",
                    "Broken reference",
                    "Формула содержит #REF!.",
                    "Расчет не может быть достоверным.",
                )
            if isinstance(cached, str) and cached.startswith("#"):
                issue(
                    issues,
                    sheet_name,
                    formula_cell.coordinate,
                    "Critical",
                    "Cached formula error",
                    f"Формула возвращает ошибку {cached}.",
                    "Расчет не может быть достоверным.",
                )


def main():
    formulas_wb = load_workbook(SOURCE, data_only=False, keep_vba=True)
    values_wb = load_workbook(SOURCE, data_only=True, keep_vba=True)
    store_names = formulas_wb.sheetnames[3:]

    monthly_rows = []
    annual_rows = []
    issues = []
    formula_counts = {}

    for store in store_names:
        fws = formulas_wb[store]
        vws = values_wb[store]
        formula_counts[store] = sum(
            1 for row in fws.iter_rows() for cell in row if is_formula(cell.value)
        )
        formulas_error_scan(fws, vws, store, issues)
        blocks = month_blocks(fws)
        if len(blocks) != 12:
            issue(
                issues,
                store,
                "A:A",
                "Critical",
                "Monthly block structure",
                f"Найдено {len(blocks)} из 12 ожидаемых месячных блоков.",
                "Годовой результат магазина неполный или непроверяемый.",
            )
        prior_data_row = None
        store_month_rows = []
        for block in blocks:
            month_record, prior_data_row = audit_month(fws, vws, store, block, issues, prior_data_row)
            monthly_rows.append(month_record)
            store_month_rows.append(month_record)

        if not store_month_rows:
            continue
        annual = {"store": store}
        for field in list(METRIC_COLUMNS) + ["net_profit"]:
            annual[field] = sum(value_or_zero(row[field]) for row in store_month_rows)
        annual["expenses"] = annual["revenue"] - annual["net_profit"]
        annual["gross_margin"] = annual["gross_profit"] / annual["revenue"] if annual["revenue"] else None
        annual["net_margin"] = annual["net_profit"] / annual["revenue"] if annual["revenue"] else None
        annual["best_month"] = max(store_month_rows, key=lambda item: item["net_profit"])["month"]
        annual["best_month_profit"] = max(store_month_rows, key=lambda item: item["net_profit"])["net_profit"]
        annual["worst_month"] = min(store_month_rows, key=lambda item: item["net_profit"])["month"]
        annual["worst_month_profit"] = min(store_month_rows, key=lambda item: item["net_profit"])["net_profit"]
        annual["loss_months"] = sum(1 for row in store_month_rows if row["net_profit"] < 0)
        annual["formula_count"] = formula_counts[store]
        annual_rows.append(annual)

    # Consolidated network monthly outputs.
    network_monthly = []
    for month_number, month in enumerate(MONTHS, start=1):
        subset = [row for row in monthly_rows if row["month_number"] == month_number]
        network_monthly.append(
            {
                "month": month,
                "month_number": month_number,
                "store_count": len(subset),
                "revenue": sum(row["revenue"] for row in subset),
                "gross_profit": sum(row["gross_profit"] for row in subset),
                "net_profit": sum(row["net_profit"] for row in subset),
                "expenses": sum(row["revenue"] - row["net_profit"] for row in subset),
            }
        )

    # Formula count variance is a useful structural red flag.
    modal_count = Counter(formula_counts.values()).most_common(1)[0][0]
    for store, count in formula_counts.items():
        if count != modal_count:
            issue(
                issues,
                store,
                "Sheet",
                "Warning",
                "Formula count variance",
                f"Число формул {count:,} отличается от модального значения {modal_count:,}.",
                "Возможна ручная вставка или удаление расчетных ячеек.",
            )

    fieldnames_monthly = [
        "store", "month", "month_number", "total_row", "net_profit_cell", *METRIC_COLUMNS.keys(), "net_profit"
    ]
    fieldnames_annual = [
        "store", *METRIC_COLUMNS.keys(), "expenses", "net_profit", "gross_margin", "net_margin",
        "best_month", "best_month_profit", "worst_month", "worst_month_profit", "loss_months", "formula_count",
    ]
    issue_fields = ["sheet", "cell", "severity", "category", "description", "impact"]
    write_csv(OUT_DIR / "store_monthly.csv", monthly_rows, fieldnames_monthly)
    write_csv(OUT_DIR / "store_annual.csv", annual_rows, fieldnames_annual)
    write_csv(OUT_DIR / "network_monthly.csv", network_monthly, ["month", "month_number", "store_count", "revenue", "gross_profit", "expenses", "net_profit"])
    write_csv(OUT_DIR / "formula_audit.csv", issues, issue_fields)

    summary = {
        "source": str(SOURCE),
        "as_of_period": "2026 calendar year (derived from filename)",
        "stores_analyzed": len(store_names),
        "store_names": store_names,
        "monthly_records": len(monthly_rows),
        "formula_counts": formula_counts,
        "formula_count_mode": modal_count,
        "issue_counts": dict(Counter(item["severity"] for item in issues)),
        "issues_total": len(issues),
        "generated_at": datetime.now().isoformat(),
        "metric_definition": {
            "revenue": "Сумма строки «Продажа» (столбец H) за месячные блоки.",
            "net_profit": "Значение формулы «Итог» (столбец AQ в первой дневной строке каждого месячного блока).",
            "expenses": "Выручка минус чистая прибыль; включает закупочную себестоимость и все статьи, вошедшие в формулу чистой прибыли.",
        },
    }
    (OUT_DIR / "audit_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    print(json.dumps({
        "stores": len(store_names),
        "monthly_records": len(monthly_rows),
        "issue_counts": summary["issue_counts"],
        "issues_total": len(issues),
        "formula_count_mode": modal_count,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
