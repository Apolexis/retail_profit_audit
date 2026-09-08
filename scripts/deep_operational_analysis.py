"""Deep operating audit for Jan–Aug 2026, excluding КА1/АЛА/МОЛ/С2/РЕЗ*.

Builds product-mix, full expense, write-off and stock analysis using cached values of
the corrected workbook. The workbook is never saved or changed.
"""
from __future__ import annotations

import json
import math
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path('/home/ubuntu/retail_profit_audit_2026')
SOURCE = Path('/home/ubuntu/upload/2026.xlsx')
OUT = ROOT/'audit_2026'/'deep_operational_excl_s2'
OUT.mkdir(parents=True, exist_ok=True)
MONTHS = ['Январь','Февраль','Март','Апрель','Май','Июнь','Июль','Август']
EXACT = {'КА1','АЛА','МОЛ','С2'}

FIELDS = {
    'purchase_smoked':'C','purchase_frozen':'D','sales_smoked':'E','sales_frozen':'F',
    'purchases':'G','revenue':'H','gross_profit':'L','cash_revenue':'M','cashless_revenue':'N',
    'receipts_total':'O','writeoff_smoked':'P','writeoff_frozen':'Q','movement':'R','discount':'S','revaluation':'T',
    'household':'U','delivery':'V','cleaning':'W','bonus':'X','service_bonus':'Y','extra_pay':'Z',
    'driver_cash':'AA','utilities_cash':'AB','other_cash_expenses':'AC','cash_operating_total':'AD',
    'driver_cashless':'AE','utilities_cashless':'AF','rent':'AG','bank_fee_raw':'AH','gross_profit_tax_raw':'AI',
    'salary_cashless':'AJ','payroll_tax':'AK','vacation_cashless':'AL','vacation_tax':'AM',
    'salary_cash':'AN','vacation_cash':'AO','personal_income_tax_raw':'AP'
}
COSTS = {
    'Хоз. нужды':'household','Доставка':'delivery','Уборка':'cleaning','Премия':'bonus','Выслуга':'service_bonus',
    'Доплата':'extra_pay','Водитель (наличные)':'driver_cash','Коммунальные (наличные)':'utilities_cash',
    'Прочие наличные расходы':'other_cash_expenses','Водитель (безнал)':'driver_cashless',
    'Коммунальные (безнал)':'utilities_cashless','Аренда':'rent','Банковская комиссия':'bank_fee',
    'Налог с валовой прибыли':'gross_profit_tax','Зарплата (безнал)':'salary_cashless','Налоги на ФОТ':'payroll_tax',
    'Отпускные (безнал)':'vacation_cashless','Налоги на отпускные':'vacation_tax','Зарплата (наличные)':'salary_cash',
    'Отпускные (наличные)':'vacation_cash','НДФЛ':'personal_income_tax','Списания М.':'writeoff_frozen'
}

def num(value):
    return float(value) if isinstance(value, (int,float)) and not isinstance(value,bool) else 0.0

def blocks(ws):
    heads=[]
    for r in range(1, ws.max_row+1):
        if ws.cell(r,1).value in MONTHS:
            heads.append((r,ws.cell(r,1).value))
    result=[]
    for i,(start,month) in enumerate(heads):
        nxt=heads[i+1][0] if i+1 < len(heads) else ws.max_row+1
        total=next((r for r in range(start+1,nxt) if ws.cell(r,1).value=='Итого'),None)
        days=[r for r in range(start+1,total or start) if isinstance(ws.cell(r,1).value,(int,float))]
        if total and days: result.append((month,total,days))
    return result

def closing_balance(ws, last_row):
    return num(ws[f'B{last_row}'].value)+num(ws[f'H{last_row}'].value)-num(ws[f'O{last_row}'].value)-num(ws[f'P{last_row}'].value)-num(ws[f'Q{last_row}'].value)+num(ws[f'R{last_row}'].value)-num(ws[f'S{last_row}'].value)+num(ws[f'T{last_row}'].value)

wb=load_workbook(SOURCE,data_only=True,read_only=False)
stores=[name for name in wb.sheetnames[3:] if name not in EXACT and not name.startswith('РЕЗ')]
monthly=[]; annual=[]
for store in stores:
    ws=wb[store]; shop=[]
    for m_index,(month,total,days) in enumerate(blocks(ws),start=1):
        rec={'store':store,'month':month,'month_number':m_index,'stock_open':num(ws[f'B{days[0]}'].value),'stock_close':closing_balance(ws,days[-1])}
        for key,col in FIELDS.items(): rec[key]=num(ws[f'{col}{total}'].value)
        for raw,clean in [('bank_fee_raw','bank_fee'),('gross_profit_tax_raw','gross_profit_tax'),('personal_income_tax_raw','personal_income_tax')]: rec[clean]=-rec[raw]
        rec['driver_total']=rec['driver_cash']+rec['driver_cashless']; rec['utilities_total']=rec['utilities_cash']+rec['utilities_cashless']
        rec['payroll_total']=sum(rec[x] for x in ['salary_cashless','payroll_tax','vacation_cashless','vacation_tax','salary_cash','vacation_cash','personal_income_tax'])
        rec['net_profit']=num(ws[f'AQ{days[0]}'].value)
        rec['frozen_sales_share']=rec['sales_frozen']/rec['revenue'] if rec['revenue'] else 0
        rec['smoked_markup']=(rec['sales_smoked']-rec['purchase_smoked'])/rec['purchase_smoked'] if rec['purchase_smoked'] else 0
        rec['frozen_markup']=(rec['sales_frozen']-rec['purchase_frozen'])/rec['purchase_frozen'] if rec['purchase_frozen'] else 0
        rec['inventory_effect']=rec['movement']-rec['discount']+rec['revaluation']
        shop.append(rec); monthly.append(rec)
    a={'store':store}
    for key in list(FIELDS)+['bank_fee','gross_profit_tax','personal_income_tax','driver_total','utilities_total','payroll_total','net_profit','inventory_effect']:
        a[key]=sum(r.get(key,0) for r in shop)
    a['stock_open']=shop[0]['stock_open']; a['stock_close']=shop[-1]['stock_close']; a['stock_change']=a['stock_close']-a['stock_open']
    a['gross_margin']=a['gross_profit']/a['revenue'] if a['revenue'] else 0; a['net_margin']=a['net_profit']/a['revenue'] if a['revenue'] else 0
    a['frozen_sales_share']=a['sales_frozen']/a['revenue'] if a['revenue'] else 0
    a['smoked_markup']=(a['sales_smoked']-a['purchase_smoked'])/a['purchase_smoked'] if a['purchase_smoked'] else 0
    a['frozen_markup']=(a['sales_frozen']-a['purchase_frozen'])/a['purchase_frozen'] if a['purchase_frozen'] else 0
    a['frozen_uplift_pp']=(a['frozen_markup']-a['smoked_markup'])*100
    a['writeoff_frozen_share']=a['writeoff_frozen']/a['revenue'] if a['revenue'] else 0
    a['writeoff_smoked_share']=a['writeoff_smoked']/a['revenue'] if a['revenue'] else 0
    a['loss_months']=sum(r['net_profit']<0 for r in shop)
    a['best_month']=max(shop,key=lambda x:x['net_profit'])['month']; a['worst_month']=min(shop,key=lambda x:x['net_profit'])['month']
    a['best_month_profit']=max(r['net_profit'] for r in shop); a['worst_month_profit']=min(r['net_profit'] for r in shop)
    annual.append(a)

adf=pd.DataFrame(annual); mdf=pd.DataFrame(monthly)
ratio_fields=list(COSTS.values())+['driver_total','utilities_total','payroll_total','writeoff_smoked']
median_ratios={field:float((adf[field]/adf['revenue']).median()) for field in ratio_fields}
median_abs={field:float(adf[field].median()) for field in ratio_fields}
for field in ratio_fields:
    adf[f'{field}_ratio']=adf[field]/adf['revenue']
    adf[f'{field}_median_gap']=adf[field]-median_ratios[field]*adf['revenue']

adf['profit_rank']=adf['net_profit'].rank(method='min',ascending=False).astype(int)
adf['margin_rank']=adf['net_margin'].rank(method='min',ascending=False).astype(int)
adf['mix_rank']=adf['frozen_sales_share'].rank(method='min',ascending=False).astype(int)
network={
    'revenue':float(adf.revenue.sum()),'gross_profit':float(adf.gross_profit.sum()),'net_profit':float(adf.net_profit.sum()),
    'net_margin':float(adf.net_profit.sum()/adf.revenue.sum()),'gross_margin':float(adf.gross_profit.sum()/adf.revenue.sum()),
    'stores':int(len(adf)),'loss_stores':int((adf.net_profit<0).sum()),'loss_amount':float(-adf.loc[adf.net_profit<0,'net_profit'].sum()),
    'frozen_markup':float((adf.sales_frozen.sum()-adf.purchase_frozen.sum())/adf.purchase_frozen.sum()),
    'smoked_markup':float((adf.sales_smoked.sum()-adf.purchase_smoked.sum())/adf.purchase_smoked.sum()),
    'frozen_sales_share':float(adf.sales_frozen.sum()/adf.revenue.sum()),
    'frozen_share_net_margin_correlation':float(adf[['frozen_sales_share','net_margin']].corr().iloc[0,1]),
    'frozen_share_gross_margin_correlation':float(adf[['frozen_sales_share','gross_margin']].corr().iloc[0,1]),
}
monthly_network=mdf.groupby(['month','month_number'],as_index=False).sum(numeric_only=True).sort_values('month_number')
monthly_network['net_margin']=monthly_network['net_profit']/monthly_network['revenue']; monthly_network['gross_margin']=monthly_network['gross_profit']/monthly_network['revenue']; monthly_network['frozen_sales_share']=monthly_network['sales_frozen']/monthly_network['revenue']
cost_summary=[]
for label,field in COSTS.items():
    cost_summary.append({'label':label,'field':field,'amount':float(adf[field].sum()),'share_revenue':float(adf[field].sum()/adf.revenue.sum()),'median_ratio':median_ratios[field],'median_abs':median_abs[field]})
cost_summary=sorted(cost_summary,key=lambda x:x['amount'],reverse=True)
data={'scope':{'period':'Январь–август 2026','stores':stores,'excluded':['КА1','АЛА','МОЛ','С2','РЕЗ*'],'source':'исправленный 2026.xlsx'},'network':network,'cost_categories':cost_summary,'median_ratios':median_ratios,'median_abs':median_abs,'stores':adf.sort_values('net_profit',ascending=False).to_dict(orient='records'),'monthly_network':monthly_network.to_dict(orient='records'),'monthly_stores':mdf.to_dict(orient='records')}
(OUT/'deep_analysis.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
adf.to_csv(OUT/'store_deep_metrics.csv',index=False,encoding='utf-8-sig'); mdf.to_csv(OUT/'monthly_deep_metrics.csv',index=False,encoding='utf-8-sig')
print(json.dumps({'stores':len(adf),'network':network,'top_costs':[x['label'] for x in cost_summary[:5]]},ensure_ascii=False,indent=2))
