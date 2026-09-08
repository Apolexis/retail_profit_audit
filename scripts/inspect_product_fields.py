from pathlib import Path
from openpyxl import load_workbook

SOURCE = Path('/home/ubuntu/upload/2026.xlsx')
OUT = Path('/home/ubuntu/retail_profit_audit_2026/audit_2026/product_field_layout.txt')

wb = load_workbook(SOURCE, data_only=False, read_only=True)
ws = wb[wb.sheetnames[3]]
lines = [f'SHEET: {ws.title}', '']
for row in range(1, 35):
    cells = []
    for col in range(1, 44):
        value = ws.cell(row, col).value
        if value is not None:
            cells.append(f'{ws.cell(row,col).coordinate}={value}')
    if cells:
        lines.append(' | '.join(cells))
OUT.write_text('\n'.join(lines), encoding='utf-8')
print(OUT)
