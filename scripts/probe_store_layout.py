"""Export a readable value/formula layout for one representative store sheet."""

from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path("/home/ubuntu/upload/учет2026.xlsm")
OUT = Path("/home/ubuntu/retail_profit_audit_2026/audit_work/store_layout_port.tsv")
SHEET_NAME = "ПОРТ"


def stringify(value):
    if value is None:
        return ""
    return str(value).replace("\t", " ").replace("\n", " ")


def main():
    formulas_wb = load_workbook(SOURCE, data_only=False, keep_vba=True)
    values_wb = load_workbook(SOURCE, data_only=True, keep_vba=True)
    fws = formulas_wb[SHEET_NAME]
    vws = values_wb[SHEET_NAME]

    with OUT.open("w", encoding="utf-8") as output:
        output.write("row\tcolumn\tvalue\tformula\n")
        for row in range(1, fws.max_row + 1):
            for col in range(1, fws.max_column + 1):
                formula_value = fws.cell(row, col).value
                cached_value = vws.cell(row, col).value
                if formula_value is not None or cached_value is not None:
                    formula = formula_value if isinstance(formula_value, str) and formula_value.startswith("=") else ""
                    display_value = cached_value if formula else formula_value
                    output.write(
                        f"{row}\t{fws.cell(row, col).coordinate}\t{stringify(display_value)}\t{stringify(formula)}\n"
                    )
    print(OUT)


if __name__ == "__main__":
    main()
