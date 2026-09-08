from pathlib import Path
import json
import openpyxl

src = Path('/home/ubuntu/upload/2026.xlsx')
out = Path('/home/ubuntu/retail_profit_audit_2026/audit_2026')
out.mkdir(exist_ok=True)

wb_values = openpyxl.load_workbook(src, data_only=True, read_only=True)
wb_formulas = openpyxl.load_workbook(src, data_only=False, read_only=True)
rows = []
for index, (value_ws, formula_ws) in enumerate(zip(wb_values.worksheets, wb_formulas.worksheets), start=1):
    formulas = 0
    month_hits = []
    for row in formula_ws.iter_rows():
        for cell in row:
            cell_value = getattr(cell, 'value', None)
            if isinstance(cell_value, str) and cell_value.startswith('='):
                formulas += 1
            if getattr(cell, 'column', None) == 1 and cell_value in {'Январь','Февраль','Март','Апрель','Май','Июнь','Июль','Август','Сентябрь','Октябрь','Ноябрь','Декабрь'}:
                month_hits.append({'month':cell_value,'row':getattr(cell, 'row', None)})
    rows.append({'index': index, 'sheet': value_ws.title, 'max_row': value_ws.max_row, 'max_col': value_ws.max_column, 'formula_cells': formulas, 'month_headers': month_hits})
summary = {'source': src.name, 'sheets_total': len(rows), 'sheets': rows, 'store_sheets': [r['sheet'] for r in rows[3:]]}
(out/'inventory.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(summary, ensure_ascii=False, indent=2))
